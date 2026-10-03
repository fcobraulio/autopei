"""Testes ponta a ponta do fluxo do PEI (requer um PostgreSQL de teste — o teste APAGA as tabelas).

    DATABASE_URL=postgresql+psycopg://autopei:autopei@localhost:5432/autopei_teste uv run pytest -q
"""
import io
import os
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("AUTOPEI_SUAP_FAKE", "1")

from docx import Document  # noqa: E402

from core import auth, fluxo  # noqa: E402
from core.db import SessionLocal, engine  # noqa: E402
from core.models import Acesso, Base, Campus, Estudante, Resposta, Usuario  # noqa: E402
from core.security import hash_senha  # noqa: E402
from core.docx_pei import gerar_docx  # noqa: E402
from core.seed import seed  # noqa: E402
from core.suap import trocar_codigo  # noqa: E402


@pytest.fixture()
def s():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as sessao:
        seed(sessao)
        sessao.commit()
        yield sessao


def _suap(s, matricula, tipo="Servidor (Docente)"):
    return auth.login_suap(s, trocar_codigo(f"fake:{matricula}:{tipo}:Pessoa {matricula}"))


def _responder(s, pei, etapa, usuario, componente=None, texto=None):
    cid = componente.id if componente else None
    for p in fluxo.perguntas(s, etapa):
        if p.tipo == "longa":
            fluxo.salvar_resposta(s, pei, p, cid, texto or "", texto is None, usuario.id)
        elif p.tipo == "checkbox":
            fluxo.salvar_resposta(s, pei, p, cid, p.opcoes[:2], False, usuario.id)
        else:
            fluxo.salvar_resposta(s, pei, p, cid, p.opcoes[0], False, usuario.id)


def _cenario(s):
    campus = s.query(Campus).one()
    s.add_all([Acesso(campus_id=campus.id, matricula="900", perfil="gestor"),
               Acesso(campus_id=campus.id, matricula="901", perfil="etep"),
               Acesso(campus_id=campus.id, matricula="psico", perfil="psicopedagogia")])
    psico = Usuario(username="psico", nome="Ana Psico", auth="local", senha_hash=hash_senha("Senha1234"))
    s.add(psico)
    s.flush()
    gestor, etep = _suap(s, "900", "Servidor (Técnico-Administrativo)"), _suap(s, "901", "Servidor")
    d1, d2, d3 = _suap(s, "111"), _suap(s, "222"), _suap(s, "333")
    sem = fluxo.criar_semestre(s, campus.id, "2026.2")
    # docente ↔ código (o 789 tem dois docentes)
    fluxo.importar_ofertas(s, sem.id, "123; Lógica; 111\n456; Banco de Dados; 222\n789; Redes; 333\n789; Redes; 111")
    est = Estudante(campus_id=campus.id, nome="Estudante Teste", matricula="2026001",
                    data_nascimento=date(2010, 5, 1), curso="Informática", nee="TEA")
    s.add(est)
    s.flush()
    return campus, sem, est, gestor, psico, etep, (d1, d2, d3)


def test_fluxo_paralelo_correcao_e_registro(s):
    campus, sem, est, gestor, psico, etep, (d1, d2, d3) = _cenario(s)
    pei = fluxo.criar_pei(s, est, sem, gestor.id, [("123", ""), ("456", ""), ("789", "")])
    assert pei.status == "psicopedagogia" and len(pei.componentes) == 3
    # docentes ainda não veem (só depois de psico e ETEP)
    assert fluxo.caixa(s, d1, {}, True) == []

    _responder(s, pei, "psicopedagogia", psico)
    # laudos: anexados pela Psicopedagogia; a gestão não anexa
    fluxo.anexar(s, pei, "psicopedagogia", psico.id, "laudo.pdf", "application/pdf", b"%PDF-1.4 teste")
    with pytest.raises(fluxo.FluxoError):
        fluxo.anexar(s, pei, "gestao", gestor.id, "outro.pdf", "application/pdf", b"%PDF")
    with pytest.raises(fluxo.FluxoError):  # psicopedagoga não pode enviar etapa da ETEP
        fluxo.enviar(s, pei, "etep", psico)
    assert fluxo.enviar(s, pei, "psicopedagogia", psico) == "etep"
    _responder(s, pei, "etep", etep)
    assert fluxo.enviar(s, pei, "etep", etep) == "docente"

    # distribuição automática e em paralelo pelos códigos
    caixa = {u.username: [c.codigo for _, _, c in fluxo.caixa(s, u, {}, True)] for u in (d1, d2, d3)}
    assert sorted(caixa["111"]) == ["123", "789"] and caixa["222"] == ["456"] and caixa["333"] == ["789"]

    c123, c456, c789 = pei.componentes
    # estudos individualizados: cadastrados pelos docentes (vários por semana, quinzenal...)
    fluxo.adicionar_estudo(s, pei, c123, d1.id, "Segunda", 13, "semanal", "Lab 3")
    fluxo.adicionar_estudo(s, pei, c123, d1.id, "Quarta", 13, "semanal")
    fluxo.adicionar_estudo(s, pei, c456, d2.id, "Sexta", 9, "quinzenal")
    with pytest.raises(fluxo.FluxoError):  # mesmo dia e hora no mesmo componente
        fluxo.adicionar_estudo(s, pei, c123, d1.id, "Segunda", 13, "mensal")
    with pytest.raises(fluxo.FluxoError):  # hora fora das horas fechadas permitidas
        fluxo.adicionar_estudo(s, pei, c123, d1.id, "Terça", 23, "semanal")
    with pytest.raises(fluxo.FluxoError):  # docente de outra disciplina não preenche
        fluxo.enviar(s, pei, "docente", d2, c123)
    for comp, doc in ((c123, d1), (c456, d2)):
        _responder(s, pei, "docente", doc, comp)
        assert fluxo.enviar(s, pei, "docente", doc, comp) == "docente"
    _responder(s, pei, "docente", d3, c789)
    assert fluxo.enviar(s, pei, "docente", d3, c789) == "revisao"
    assert fluxo.caixa(s, d1, {}, True) == []  # 789 já concluído por d3

    # correção pontual na ETEP: volta só para a ETEP e depois direto à revisão
    fluxo.devolver(s, pei, gestor.id, "etep", "Ajustar histórico")
    assert pei.status == "etep" and pei.em_correcao
    assert fluxo.enviar(s, pei, "etep", etep) == "revisao"
    assert all(c.status == "concluido" for c in pei.componentes)

    # correção em cascata (explícita): ETEP e depois todos os docentes de novo
    fluxo.devolver(s, pei, gestor.id, "etep", "Refazer com os docentes", cascata=True)
    assert fluxo.enviar(s, pei, "etep", etep) == "docente"
    assert all(c.status == "pendente" for c in pei.componentes)
    for comp, doc in ((c123, d1), (c456, d2), (c789, d1)):
        fluxo.enviar(s, pei, "docente", doc, comp)
    assert pei.status == "revisao"

    # correção de um docente só
    fluxo.devolver(s, pei, gestor.id, "docente", "Detalhar avaliação", c456.id)
    assert [c.codigo for _, _, c in fluxo.caixa(s, d2, {}, True)] == ["456"]
    assert fluxo.caixa(s, d1, {}, True) == []
    assert fluxo.enviar(s, pei, "docente", d2, c456) == "revisao"

    # anexos são internos
    assert len(fluxo.anexos(s, pei.id)) == 1
    with pytest.raises(fluxo.FluxoError):  # componente concluído: estudos travados
        fluxo.adicionar_estudo(s, pei, c123, d1.id, "Terça", 14, "semanal")

    with pytest.raises(fluxo.FluxoError):
        fluxo.aprovar(s, pei, gestor.id)  # parecer da equipe em branco
    _responder(s, pei, "gestao", gestor)
    fluxo.aprovar(s, pei, gestor.id)
    assert pei.status == "aprovado" and pei.docx
    with pytest.raises(fluxo.FluxoError):
        fluxo.encerrar_semestre(s, sem, gestor.id)
    fluxo.registrar(s, pei, gestor.id)
    fluxo.encerrar_semestre(s, sem, gestor.id)
    s.commit()

    d = Document(io.BytesIO(pei.docx))
    assert not d.inline_shapes  # sem imagens
    celulas = " ".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "ESTUDANTE TESTE" in celulas and "456 — Banco de Dados" in celulas
    assert "laudo.pdf" not in celulas  # anexos não vão para o DOCX
    assert "não precisaram ser adaptados" in celulas  # "nada a declarar"
    estudos = next(t for t in d.tables if t.rows[0].cells[0].text == "Componente curricular"
                   and len(t.columns) == 6)
    linhas = [[c.text for c in r.cells] for r in estudos.rows[1:]]
    assert linhas[0][2:5] == ["Segunda", "13h–14h", "Semanal (toda semana)"]
    assert len(linhas) == 3 and linhas[2][4].startswith("Quinzenal")


def test_pre_preenchimento_do_semestre_anterior(s):
    campus, sem, est, gestor, psico, etep, (d1, d2, d3) = _cenario(s)
    pei = fluxo.criar_pei(s, est, sem, gestor.id, [("123", ""), ("456", "")])
    _responder(s, pei, "psicopedagogia", psico, texto="Texto da psicopedagoga")
    fluxo.enviar(s, pei, "psicopedagogia", psico)
    _responder(s, pei, "etep", etep, texto="Texto da ETEP")
    fluxo.enviar(s, pei, "etep", etep)
    _responder(s, pei, "docente", d1, pei.componentes[0], texto="Lógica em 2026.2")
    s.flush()

    sem2 = fluxo.criar_semestre(s, campus.id, "2027.1")
    fluxo.salvar_oferta(s, sem2.id, "123", "Lógica", "111")
    novo = fluxo.criar_pei(s, est, sem2, gestor.id, [("123", ""), ("999", "Nova disciplina")])
    assert novo.prefill_origem == "2026.2"
    geral = fluxo.respostas(s, novo.id)
    assert geral and all(r.origem == "2026.2" for r in geral.values())
    assert any((r.valor or {}).get("v") == "Texto da ETEP" for r in geral.values())
    c123, c999 = novo.componentes
    assert any((r.valor or {}).get("v") == "Lógica em 2026.2" for r in fluxo.respostas(s, novo.id, c123.id).values())
    assert fluxo.respostas(s, novo.id, c999.id) == {}  # disciplina nova começa em branco
    # alterar uma resposta remove a marca de pré-preenchimento
    p = fluxo.perguntas(s, "etep")[0]
    fluxo.salvar_resposta(s, novo, p, None, "Atualizado em 2027.1", False, etep.id)
    s.flush()
    r = s.query(Resposta).filter_by(pei_id=novo.id, pergunta_id=p.id, componente_id=None).one()
    assert r.origem == "" and r.valor["v"] == "Atualizado em 2027.1"
    # nenhum docente cadastrou estudo individualizado → o DOCX informa
    texto = " ".join(par.text for par in Document(io.BytesIO(gerar_docx(s, novo))).paragraphs)
    assert "não tem nenhum estudo individualizado cadastrado" in texto


def test_regras_de_login(s):
    campus = s.query(Campus).one()
    s.add(Acesso(campus_id=campus.id, matricula="psico", perfil="psicopedagogia"))
    s.add(Usuario(username="psico", nome="Ana", auth="local", senha_hash=hash_senha("Senha1234")))
    s.add(Usuario(username="outro", nome="Outro", auth="local", senha_hash=hash_senha("Senha1234")))
    s.add(Acesso(campus_id=campus.id, matricula="777", perfil="etep"))
    s.flush()
    assert auth.login_local(s, "psico", "Senha1234").username == "psico"
    with pytest.raises(auth.LoginError):  # conta local sem perfil de psicopedagogia
        auth.login_local(s, "outro", "Senha1234")
    u = _suap(s, "777", "Servidor (Técnico-Administrativo)")
    ctx = auth.montar_contexto(s, u.id, "suap", None)
    assert ctx.perfis == {"etep"} and ctx.docente
    # qualquer pessoa do SUAP entra (como docente), mesmo sem estar na lista
    v = _suap(s, "555")
    ctx = auth.montar_contexto(s, v.id, "suap", None)
    assert ctx.perfis == set() and ctx.docente and ctx.campus_id is None


def test_opcao_outra_pede_descricao_e_vai_para_o_docx(s):
    campus, sem, est, gestor, psico, etep, docentes = _cenario(s)
    pei = fluxo.criar_pei(s, est, sem, gestor.id, [("123", "")])
    _responder(s, pei, "psicopedagogia", psico)
    nee = next(p for p in fluxo.perguntas(s, "psicopedagogia") if fluxo.opcao_outra(p))
    outra = fluxo.opcao_outra(nee)
    # marcou "Outra" sem descrever: não deixa enviar
    fluxo.salvar_resposta(s, pei, nee, None, [nee.opcoes[0], outra], False, psico.id)
    s.flush()
    assert any("descreva" in f for f in fluxo.faltantes(s, pei, "psicopedagogia"))
    with pytest.raises(fluxo.FluxoError):
        fluxo.enviar(s, pei, "psicopedagogia", psico)
    # com a descrição, envia e o DOCX mostra "Outra: <descrição>"
    fluxo.salvar_resposta(s, pei, nee, None, [nee.opcoes[0], outra], False, psico.id,
                          "Transtorno do processamento auditivo central")
    s.flush()
    assert fluxo.faltantes(s, pei, "psicopedagogia") == []
    assert fluxo.enviar(s, pei, "psicopedagogia", psico) == "etep"
    texto = " ".join(c.text for t in Document(io.BytesIO(gerar_docx(s, pei))).tables
                     for r in t.rows for c in r.cells)
    assert "Outra: Transtorno do processamento auditivo central" in texto
    assert "descrever abaixo" not in texto
