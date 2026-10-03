"""Geração do DOCX do PEI.

O documento é feito para ser copiado e colado no editor do SUAP, portanto usa
apenas texto (negrito, itálico, sublinhado, cores, maiúsculas/minúsculas) e
tabelas. Não há imagens, cabeçalhos/rodapés de página, caixas de texto nem
campos automáticos.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.orm import Session

from . import config
from .fluxo import docentes_do_componente, estudos, perguntas, respostas, rotulo_componente
from .models import FREQUENCIAS, Pei, Pergunta, Resposta, Usuario, faixa_hora

VERDE = RGBColor(0x1B, 0x5E, 0x20)
VERMELHO = RGBColor(0xC6, 0x28, 0x28)
CINZA = RGBColor(0x55, 0x55, 0x55)
FUNDO_TITULO = "D8D8D8"
FUNDO_COMPONENTE = "E2EFD9"
FONTE = "Arial"


# ---------------------------------------------------------------------------
# Auxiliares de formatação
# ---------------------------------------------------------------------------
def _sombrear(cell, cor: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tc_pr.append(shd)


def _run(par, texto: str, negrito=False, italico=False, sublinhado=False, cor=None, tam=None):
    r = par.add_run(texto)
    r.bold, r.italic, r.underline = negrito, italico, sublinhado
    if cor is not None:
        r.font.color.rgb = cor
    if tam:
        r.font.size = Pt(tam)
    return r


def _par(container, texto: str = "", alinhar=None, espaco_depois=4, **fmt):
    par = container.add_paragraph()
    par.paragraph_format.space_after = Pt(espaco_depois)
    par.paragraph_format.space_before = Pt(0)
    if alinhar is not None:
        par.alignment = alinhar
    if texto:
        _run(par, texto, **fmt)
    return par


def _celula_texto(cell, texto: str = "", **fmt):
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    if texto:
        _run(par, texto, **fmt)
    return par


def _titulo_secao(doc, numero: int, texto: str) -> None:
    par = _par(doc, espaco_depois=6)
    par.paragraph_format.space_before = Pt(12)
    _run(par, f"{numero}. {texto.upper()}", negrito=True, cor=VERDE, tam=11)


def _tabela(doc, linhas: int, colunas: int):
    t = doc.add_table(rows=linhas, cols=colunas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    return t


def _idade(nasc: date | None) -> str:
    if not nasc:
        return "não informada"
    hoje = date.today()
    anos = hoje.year - nasc.year - ((hoje.month, hoje.day) < (nasc.month, nasc.day))
    return f"{anos} anos"


def _valor_texto(p: Pergunta, r: Resposta | None) -> tuple[list[str], bool]:
    """Retorna (linhas, automatico)."""
    if r is None:
        return [], False
    if r.nada_declarar:
        return [p.texto_nada or "Nada a declarar."], True
    v = (r.valor or {}).get("v")
    if isinstance(v, list):
        return [f"• {item}" for item in v], False
    texto = str(v or "").strip()
    return ([linha for linha in texto.splitlines() if linha.strip()] if texto else []), False


def _bloco_perguntas(doc, lista: list[Pergunta], resp: dict[int, Resposta]) -> None:
    """Uma tabela de uma coluna por seção: título sombreado + uma linha por pergunta."""
    secoes: dict[str, list[Pergunta]] = {}
    for p in lista:
        secoes.setdefault(p.secao or "Informações", []).append(p)
    for secao, itens in secoes.items():
        t = _tabela(doc, 1 + len(itens), 1)
        cab = t.rows[0].cells[0]
        _sombrear(cab, FUNDO_TITULO)
        _celula_texto(cab, secao.upper(), negrito=True).paragraph_format.keep_with_next = True
        for i, p in enumerate(itens, start=1):
            cell = t.rows[i].cells[0]
            linhas, auto = _valor_texto(p, resp.get(p.id))
            par = _celula_texto(cell)
            par.paragraph_format.keep_with_next = True
            _run(par, p.enunciado, negrito=True)
            if p.tipo in ("radio", "select") and linhas:
                _run(par, " " if p.enunciado.rstrip()[-1:] in "?:" else ": ")
                _run(par, linhas[0])
                continue
            if not linhas:
                _par(cell, "Não informado.", italico=True, cor=CINZA, espaco_depois=2)
                continue
            for linha in linhas:
                _par(cell, linha, italico=auto, espaco_depois=2,
                     alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY if not auto and len(linha) > 90 else None)
        _par(doc, espaco_depois=6)


def _responsaveis_etapa(pei: Pei, termo: str) -> list[str]:
    nomes = []
    for h in pei.historico:
        if termo in h.acao and h.usuario and h.usuario.nome not in nomes:
            nomes.append(h.usuario.nome)
    return nomes


# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------
def nome_arquivo(pei: Pei) -> str:
    base = unicodedata.normalize("NFKD", pei.estudante.nome).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_")
    return f"PEI_{base}_{pei.semestre.codigo}.docx"


def gerar_docx(s: Session, pei: Pei) -> bytes:
    est = pei.estudante
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = sec.right_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = FONTE
    estilo.font.size = Pt(10)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)

    # --- Cabeçalho (texto) ---------------------------------------------------
    centro = WD_ALIGN_PARAGRAPH.CENTER
    _par(doc, config.INSTITUICAO, alinhar=centro, negrito=True, espaco_depois=0)
    _par(doc, f"CAMPUS {pei.campus.nome.upper()}", alinhar=centro, negrito=True, espaco_depois=0)
    _par(doc, "DIRETORIA ACADÊMICA (DIAC)", alinhar=centro, negrito=True, espaco_depois=0)
    _par(doc, "NÚCLEO DE APOIO ÀS PESSOAS COM NECESSIDADES EDUCACIONAIS ESPECÍFICAS (NAPNE)",
         alinhar=centro, negrito=True, espaco_depois=10)
    _par(doc, "PLANO EDUCACIONAL INDIVIDUALIZADO (PEI)", alinhar=centro, negrito=True,
         cor=VERDE, tam=14, espaco_depois=0)
    _par(doc, f"Semestre letivo {pei.semestre.codigo}", alinhar=centro, italico=True, espaco_depois=10)

    alerta = _par(doc, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY, espaco_depois=8)
    _run(alerta, "Alerta Ético:", negrito=True, sublinhado=True, cor=VERMELHO)
    _run(alerta, " as informações contidas neste documento são consideradas reservadas e o "
                 "compartilhamento das mesmas deve ser restrito apenas às/aos envolvidos na ação "
                 "pedagógica, sob pena de implicações legais.")

    n = 1
    # --- 1. Identificação ----------------------------------------------------
    _titulo_secao(doc, n, "Identificação do(a) estudante")
    psico = _responsaveis_etapa(pei, "Psicopedagogia concluiu")
    etep = _responsaveis_etapa(pei, "ETEP concluiu")
    aprov = s.get(Usuario, pei.aprovado_por_id) if pei.aprovado_por_id else None
    equipe = []
    if psico:
        equipe.append("Psicopedagogia: " + ", ".join(psico))
    if etep:
        equipe.append("ETEP: " + ", ".join(etep))
    if aprov:
        equipe.append("NAPNE: " + aprov.nome)
    nasc = est.data_nascimento.strftime("%d/%m/%Y") if est.data_nascimento else "não informada"
    linhas = [
        ("Nome do(a) estudante", est.nome.upper()),
        ("Matrícula", est.matricula or "—"),
        ("Data de nascimento / Idade", f"{nasc} ({_idade(est.data_nascimento)})" if est.data_nascimento else "não informada"),
        ("Curso", est.curso or "—"),
        ("Nível de ensino / Forma", est.nivel_forma or "—"),
        ("Ano / Período", est.periodo or "—"),
        ("Turno", est.turno or "—"),
        ("Necessidades Educacionais Específicas", est.nee or "—"),
        ("Responsável legal", est.responsavel or "—"),
        ("Equipe multiprofissional responsável", "; ".join(equipe) or "ETEP / NAPNE"),
    ]
    t = _tabela(doc, len(linhas), 2)
    for i, (rot, val) in enumerate(linhas):
        _celula_texto(t.rows[i].cells[0], rot, negrito=True)
        _sombrear(t.rows[i].cells[0], FUNDO_TITULO)
        _celula_texto(t.rows[i].cells[1], val, negrito=(i == 0))
        t.rows[i].cells[0].width = Cm(6)
        t.rows[i].cells[1].width = Cm(11)

    # --- 2. Componentes e docentes ---------------------------------------------
    n += 1
    _titulo_secao(doc, n, "Componentes curriculares cursados no período")
    comps = pei.componentes
    rotulo = {c.id: rotulo_componente(s, c) for c in comps}
    docentes = {}
    for c in comps:
        nomes = docentes_do_componente(s, c)
        if c.concluido_por and c.concluido_por.nome not in nomes:
            nomes.insert(0, c.concluido_por.nome)
        docentes[c.id] = ", ".join(nomes) or "—"
    t = _tabela(doc, 1 + max(1, len(comps)), 2)
    for j, rot in enumerate(("Componente curricular", "Docente(s)")):
        _sombrear(t.rows[0].cells[j], FUNDO_TITULO)
        _celula_texto(t.rows[0].cells[j], rot, negrito=True)
    if comps:
        for i, c in enumerate(comps, start=1):
            _celula_texto(t.rows[i].cells[0], rotulo[c.id])
            _celula_texto(t.rows[i].cells[1], docentes[c.id])
    else:
        _celula_texto(t.rows[1].cells[0], "Nenhum componente informado.", italico=True)

    # --- 3. Estudos individualizados (cadastrados pelos docentes) -------------------
    n += 1
    _titulo_secao(doc, n, "Estudos individualizados")
    lista = estudos(s, pei.id)
    if not lista:
        _par(doc, "Este(a) estudante não tem nenhum estudo individualizado cadastrado pelos docentes "
                  "neste período.", italico=True)
    else:
        _par(doc, "Horários de estudo individualizado oferecidos pelos docentes. Cada linha é um horário; "
                  "um mesmo componente pode ter mais de um horário.", tam=9, cor=CINZA)
        cab = ("Componente curricular", "Docente", "Dia", "Horário", "Frequência", "Local / observação")
        larg = (4.0, 3.0, 1.9, 2.0, 3.3, 3.0)
        t = _tabela(doc, 1 + len(lista), len(cab))
        for j, rot in enumerate(cab):
            _sombrear(t.rows[0].cells[j], FUNDO_TITULO)
            _celula_texto(t.rows[0].cells[j], rot, negrito=True)
        comp_por_id = {c.id: c for c in comps}
        for i, e in enumerate(lista, start=1):
            c = comp_por_id.get(e.componente_id)
            quem = e.usuario.nome if e.usuario else (docentes[c.id] if c else "—")
            valores = (rotulo[c.id] if c else "—", quem, e.dia, faixa_hora(e.hora),
                       FREQUENCIAS.get(e.frequencia, e.frequencia), e.local or "—")
            for j, v in enumerate(valores):
                _celula_texto(t.rows[i].cells[j], v)
        t.autofit = False
        for j, w in enumerate(larg):
            t.columns[j].width = Cm(w)
            for cell in t.columns[j].cells:
                cell.width = Cm(w)

    # --- 4. Psicopedagogia -----------------------------------------------------
    resp_geral = respostas(s, pei.id)
    n += 1
    _titulo_secao(doc, n, "Avaliação psicopedagógica")
    _bloco_perguntas(doc, perguntas(s, "psicopedagogia", pei.id), resp_geral)

    # --- 5. ETEP ---------------------------------------------------------------
    n += 1
    _titulo_secao(doc, n, "Histórico e perfil pedagógico (ETEP)")
    _bloco_perguntas(doc, perguntas(s, "etep", pei.id), resp_geral)

    # --- 6. Componentes --------------------------------------------------------
    n += 1
    _titulo_secao(doc, n, "Adaptações curriculares por componente")
    p_doc = perguntas(s, "docente", pei.id)
    if not comps:
        _par(doc, "Não há componentes curriculares vinculados a este PEI.", italico=True)
    for c in comps:
        t = _tabela(doc, 2, 1)
        _sombrear(t.rows[0].cells[0], FUNDO_COMPONENTE)
        par = _celula_texto(t.rows[0].cells[0])
        _run(par, "Componente: ", negrito=True)
        _run(par, rotulo[c.id].upper(), negrito=True, cor=VERDE)
        _sombrear(t.rows[1].cells[0], FUNDO_COMPONENTE)
        par = _celula_texto(t.rows[1].cells[0])
        _run(par, "Docente(s): ", negrito=True)
        _run(par, docentes[c.id])
        _par(doc, espaco_depois=4)
        _bloco_perguntas(doc, p_doc, respostas(s, pei.id, c.id))

    # --- 7. Registro e parecer da equipe ------------------------------------------
    n += 1
    _titulo_secao(doc, n, "Registro do acompanhamento do PEI e parecer da equipe")
    _bloco_perguntas(doc, perguntas(s, "gestao", pei.id), resp_geral)

    # --- 8. Termo de ciência ---------------------------------------------------------
    n += 1
    _titulo_secao(doc, n, "Termo de ciência")
    just = WD_ALIGN_PARAGRAPH.JUSTIFY
    _par(doc, "Eu, responsável pelo(a) estudante, recebi as informações e orientações a respeito "
              "deste Plano Educacional Individualizado (PEI) para melhor atender às necessidades do(a) "
              "estudante e, assim, promover o seu desenvolvimento relacionado à aprendizagem.", alinhar=just)
    _par(doc, "Como família, comprometo-me a atualizar as informações médicas do(a) estudante, bem como "
              "dos especialistas que o(a) acompanham, e a informar a instituição sobre as necessidades "
              "específicas deste(a). É imprescindível que a família acompanhe e promova o desenvolvimento "
              "do(a) estudante em suas áreas cognitivas, emocionais e sociais, firmando assim uma parceria "
              "com a instituição.", alinhar=just)
    _par(doc, "Cientes, assinam eletronicamente no SUAP:", negrito=True, espaco_depois=6)

    envolvidos = [("Responsável legal", est.responsavel or "")]
    for c in comps:
        envolvidos.append((f"Docente – {rotulo[c.id]}", docentes[c.id]))
    envolvidos.append(("Coordenação de Curso", ""))
    envolvidos.append(("NAPNE (responsável)", aprov.nome if aprov else ""))
    envolvidos.append(("ETEP (responsável)", ", ".join(etep)))
    envolvidos.append(("Psicopedagogia", ", ".join(psico)))
    envolvidos.append(("Outros profissionais envolvidos", ""))
    t = _tabela(doc, 1 + len(envolvidos), 2)
    for j, rot in enumerate(("Função", "Nome")):
        _sombrear(t.rows[0].cells[j], FUNDO_TITULO)
        _celula_texto(t.rows[0].cells[j], rot, negrito=True)
    for i, (funcao, nome) in enumerate(envolvidos, start=1):
        _celula_texto(t.rows[i].cells[0], funcao, negrito=True)
        _celula_texto(t.rows[i].cells[1], nome)

    data_ref = pei.aprovado_em or datetime.now()
    _par(doc, espaco_depois=6)
    _par(doc, f"{pei.campus.cidade or pei.campus.nome}/RN, {data_ref.strftime('%d/%m/%Y')}.",
         alinhar=WD_ALIGN_PARAGRAPH.RIGHT)
    _par(doc, f"Documento gerado pelo AutoPEI em {datetime.now().strftime('%d/%m/%Y %H:%M')}.",
         italico=True, cor=CINZA, tam=8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
